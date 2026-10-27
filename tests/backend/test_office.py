"""PowerPoint, Word and Excel builders (backend/fsa/office.py): structure, content, validation, real readers."""
import os
import shutil
import subprocess
import tempfile
import unittest
import zipfile
from xml.etree import ElementTree as ET

import helpers  # noqa: F401  (puts the repo on sys.path)

from backend.fsa import office
from backend.fsa.store import Invalid

SOURCES = ["SMS Manual Rev 5, §3.2, p.14", "Audit Handbook Rev 2, §4.1, p.22"]
PPTX = {"subtitle": "Safety Review Board", "slides": [
    {"title": "Key findings", "bullets": ["Unstable approaches up **18%**", {"text": "A320 night arrivals", "level": 1}],
     "note": "Source: FDM Manual §5"},
    {"title": "Exceedances", "table": {"columns": ["Event", "Count"], "rows": [["Unstable approach", 37], ["GPWS", 3]]}},
    {"title": "Actions", "bullets": ["Brief crews"], "table": {"columns": ["Action", "Owner"], "rows": [["Brief", "Ops"]]}}]}
DOCX = {"subtitle": "WR-1", "sections": [
    {"heading": "Summary", "paragraphs": ["Meets **most** of §3.", "Two gaps."]},
    {"heading": "Findings", "level": 2, "bullets": ["No monthly review", {"text": "Last in June", "level": 1}]},
    {"heading": "Actions", "table": {"columns": ["#", "Action"], "rows": [[1, "Monthly review"]]}}]}
XLSX = {"sheets": [{"name": "Findings: Q3/2026", "columns": ["ID", "Finding", "Score"],
                    "rows": [["F-1", "No monthly review", 3], ["F-2", "No trend <analysis> & more", 1.5]]},
                   {"name": "Findings: Q3/2026", "columns": ["x"], "rows": []}]}


def parts(path):
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n) for n in z.namelist()}


class Base(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)

    def build(self, fmt, content, sources=SOURCES, title="Q3 Review"):
        path = os.path.join(self.tmp.name, f"{fmt}-file.{fmt}")
        info = office.build(path, fmt, title, content, sources)
        return path, info

    def check_package(self, path):
        p = parts(path)
        self.assertEqual(list(p)[0], "[Content_Types].xml")
        ct = ET.fromstring(p["[Content_Types].xml"])
        overrides = {o.get("PartName") for o in ct if o.tag.endswith("Override")}
        for name, data in p.items():
            if name.endswith((".xml", ".rels")):
                ET.fromstring(data)  # well-formed
            if name.endswith(".xml") and name != "[Content_Types].xml":
                self.assertIn("/" + name, overrides, name)
        for name, data in p.items():  # every relationship points at a part that exists
            if name.endswith(".rels"):
                base = os.path.dirname(os.path.dirname(name))
                for rel in ET.fromstring(data):
                    target = os.path.normpath(os.path.join(base, rel.get("Target"))).replace(os.sep, "/")
                    self.assertIn(target, p, f"{name} -> {rel.get('Target')}")
        return p


class BuildTests(Base):
    def test_pptx(self):
        path, info = self.build("pptx", PPTX)
        self.assertEqual(info["slides"], 5)  # title + 3 + sources
        p = self.check_package(path)
        s2 = p["ppt/slides/slide2.xml"].decode()
        self.assertIn("Key findings", s2)
        self.assertIn('b="1" dirty="0"><a:solidFill><a:srgbClr val="1B1F27"/></a:solidFill><a:latin typeface="Calibri"/>'
                      '</a:rPr><a:t>18%', s2)  # **bold**
        self.assertIn('lvl="1"', s2)
        self.assertIn("<a:tbl>", p["ppt/slides/slide3.xml"].decode())
        self.assertIn("Audit Handbook", p["ppt/slides/slide5.xml"].decode())
        self.assertIn("Q3 Review", p["ppt/slides/slide1.xml"].decode())

    def test_docx(self):
        path, info = self.build("docx", DOCX)
        p = self.check_package(path)
        doc = p["word/document.xml"].decode()
        for s in ('w:val="Title"', 'w:val="Heading2"', 'w:val="ListBullet"', "<w:tbl>", "Sources", "<w:b/>"):
            self.assertIn(s, doc)

    def test_xlsx(self):
        path, info = self.build("xlsx", XLSX)
        self.assertEqual(info["sheets"], ["Findings  Q3 2026", "Findings  Q3 2026 (2)", "Sources"])
        p = self.check_package(path)
        sh = p["xl/worksheets/sheet1.xml"].decode()
        self.assertIn("<v>1.5</v>", sh)  # numbers stay numbers
        self.assertIn("&lt;analysis&gt; &amp; more", sh)  # text is escaped

    def test_validation(self):
        for fmt, bad, msg in [
            ("pptx", {}, "slides"), ("pptx", {"slides": [{"bullets": ["x"]}]}, "title"),
            ("pptx", {"slides": [{"title": "t", "bullets": ["b"] * 13}]}, "12"),
            ("pptx", {"slides": [{"title": "t", "table": {"rows": [[1]]}}]}, "columns"),
            ("docx", {"sections": "x"}, "sections"), ("xlsx", {"sheets": [{"name": "a"}]}, "columns"),
            ("csv", {}, "format"),
        ]:
            with self.assertRaises(Invalid) as e:
                self.build(fmt, bad)
            self.assertIn(msg, str(e.exception), fmt)
        with self.assertRaises(Invalid):
            office.build(os.path.join(self.tmp.name, "x.pptx"), "pptx", " ", PPTX)

    def test_control_characters_are_dropped(self):
        path, _ = self.build("docx", {"sections": [{"heading": "A\x00B\x07", "paragraphs": ["ok\x1b"]}]})
        ET.fromstring(parts(path)["word/document.xml"])


def _libs():
    try:
        import docx  # noqa: F401
        import openpyxl  # noqa: F401
        import pptx  # noqa: F401
        return True
    except ImportError:
        return False


@unittest.skipUnless(_libs(), "python-pptx, python-docx and openpyxl not installed (test-only)")
class ReaderTests(Base):
    """Independent readers open the files and see the content."""

    def test_python_pptx(self):
        import pptx
        path, _ = self.build("pptx", PPTX)
        pres = pptx.Presentation(path)
        self.assertEqual(len(pres.slides), 5)
        texts = [sh.text_frame.text for sh in pres.slides[1].shapes if sh.has_text_frame]
        self.assertTrue(any("Unstable approaches up 18%" in t for t in texts))
        tables = [sh for sh in pres.slides[2].shapes if sh.has_table]
        self.assertEqual(tables[0].table.cell(1, 1).text, "37")

    def test_python_docx(self):
        import docx
        path, _ = self.build("docx", DOCX)
        d = docx.Document(path)
        styles = [(p.style.name, p.text) for p in d.paragraphs]
        self.assertIn(("Title", "Q3 Review"), styles)
        self.assertIn(("Heading 2", "Findings"), styles)
        self.assertEqual(d.tables[0].cell(1, 1).text, "Monthly review")

    def test_openpyxl(self):
        import openpyxl
        path, _ = self.build("xlsx", XLSX)
        wb = openpyxl.load_workbook(path)
        ws = wb["Findings  Q3 2026"]
        self.assertEqual((ws["A1"].value, ws["C3"].value, ws["A1"].font.b), ("ID", 1.5, True))
        self.assertEqual(wb["Sources"]["A2"].value, SOURCES[0])


@unittest.skipUnless(shutil.which("soffice"), "LibreOffice not installed")
class LibreOfficeTests(Base):
    def test_all_three_convert_to_pdf(self):
        files = [self.build(fmt, c)[0] for fmt, c in (("pptx", PPTX), ("docx", DOCX), ("xlsx", XLSX))]
        out = os.path.join(self.tmp.name, "pdf")
        env = dict(os.environ, HOME=self.tmp.name)  # a private LibreOffice profile
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", out, *files],
                       check=True, capture_output=True, timeout=180, env=env)
        for f in files:
            pdf = os.path.join(out, os.path.splitext(os.path.basename(f))[0] + ".pdf")
            self.assertGreater(os.path.getsize(pdf), 5000, pdf)


if __name__ == "__main__":
    unittest.main()
